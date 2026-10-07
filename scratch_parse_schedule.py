import openpyxl
import sys
import json

sys.stdout.reconfigure(encoding='utf-8')
path = r'C:\Users\SHREE\Desktop\SCHEDULE VS DISPATCH October FINAL v6.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ms = wb['Master Sheet']

ai_summary = {}
rows_parsed = 0

for r in range(22, ms.max_row + 1):
    raw_ai = ms.cell(r, 5).value
    if raw_ai is None or str(raw_ai).strip() == '':
        continue
    
    ai_str = str(raw_ai).strip()
    if ai_str.endswith('.0'):
        ai_str = ai_str[:-2]
    norm_ai = ai_str.upper() if ai_str.upper().startswith('AI-') else f'AI-{ai_str}'
    
    sch_val = ms.cell(r, 14).value
    try:
        sch_qty = float(sch_val) if sch_val is not None else 0.0
    except:
        sch_qty = 0.0
        
    cust = str(ms.cell(r, 1).value or '').strip()
    part_no = str(ms.cell(r, 2).value or '').strip()
    part_name = str(ms.cell(r, 6).value or '').strip()
    c7 = str(ms.cell(r, 7).value or '').strip()
    c8 = str(ms.cell(r, 8).value or '').strip()
    c10 = str(ms.cell(r, 10).value or '').strip()
    
    # Classify machine category
    machine_cat = 'Sliding Head'
    if 'AUTO' in c7 and 'NON' not in c7:
        machine_cat = 'Sliding Head'
    elif 'NON AUTO' in c7:
        p_lower = part_name.lower()
        if 'grind' in p_lower:
            machine_cat = 'Grinding'
        elif 'roll' in p_lower:
            machine_cat = 'Rolling'
        elif c8 == 'VMC' or c10 == 'MILLING' or 'milling' in p_lower:
            machine_cat = 'VMC / Milling'
        else:
            machine_cat = 'CNC'
    else:
        machine_cat = 'Sliding Head'
    
    if norm_ai not in ai_summary:
        ai_summary[norm_ai] = {
            'aiNumber': norm_ai,
            'totalSchedule': 0,
            'machineCategory': machine_cat,
            'partName': part_name,
            'partNo': part_no,
            'customers': []
        }
    
    ai_summary[norm_ai]['totalSchedule'] += int(sch_qty)
    if cust and cust not in ai_summary[norm_ai]['customers']:
        ai_summary[norm_ai]['customers'].append(cust)
    if part_name and not ai_summary[norm_ai]['partName']:
        ai_summary[norm_ai]['partName'] = part_name
    if part_no and not ai_summary[norm_ai]['partNo']:
        ai_summary[norm_ai]['partNo'] = part_no
    rows_parsed += 1

out_list = []
for k in sorted(ai_summary.keys(), key=lambda x: int(x.replace('AI-', '')) if x.replace('AI-', '').isdigit() else 99999):
    v = ai_summary[k]
    out_list.append({
        'aiNumber': v['aiNumber'],
        'totalSchedule': v['totalSchedule'],
        'machineCategory': v['machineCategory'],
        'partName': v['partName'] or 'Precision Turned Part',
        'partNo': v['partNo'] or '-'
    })

with open(r'c:\Users\SHREE\Desktop\DIGITALIZATION\src\data\octoberScheduleData.json', 'w', encoding='utf-8') as f:
    json.dump(out_list, f, indent=2)

print(f"Total AI Numbers: {len(out_list)}")
sliding_count = sum(1 for x in out_list if x['machineCategory'] == 'Sliding Head')
cnc_count = sum(1 for x in out_list if x['machineCategory'] == 'CNC')
grinding_count = sum(1 for x in out_list if x['machineCategory'] == 'Grinding')
rolling_count = sum(1 for x in out_list if x['machineCategory'] == 'Rolling')
vmc_count = sum(1 for x in out_list if x['machineCategory'] == 'VMC / Milling')

print(f"Sliding Head: {sliding_count}")
print(f"CNC: {cnc_count}")
print(f"Grinding: {grinding_count}")
print(f"Rolling: {rolling_count}")
print(f"VMC/Milling: {vmc_count}")
